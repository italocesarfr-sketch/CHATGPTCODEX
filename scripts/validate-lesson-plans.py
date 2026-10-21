"""Leitura independente dos planos produzidos pelo escritor Kotlin, sem regravar XLSX."""
from pathlib import Path
from datetime import datetime
import json, sys, zipfile
import openpyxl
root = Path(sys.argv[1] if len(sys.argv)>1 else 'app/build/plans-export-validation')
results=[]
for name in ('single','plans'):
    expected=json.loads((root/f'{name}.json').read_text())
    with zipfile.ZipFile(root/f'{name}.xlsx') as archive:
        assert archive.testzip() is None
    workbook=openpyxl.load_workbook(root/f'{name}.xlsx',data_only=False)
    assert len(workbook.worksheets)==1
    sheet=workbook.active
    assert sheet.max_column==3 and sheet.max_row==len(expected)+2
    assert [sheet.cell(2,c).value for c in range(1,4)]==['Data','Título','Descrição']
    assert not sheet.protection.sheet and sheet.freeze_panes=='A3'
    for row, source in enumerate(expected,3):
        date,title,description=[sheet.cell(row,c) for c in range(1,4)]
        assert isinstance(date.value,datetime) and date.value.date().isoformat()==source['date']
        assert date.number_format=='dd/mm/yyyy'
        assert title.value==source['title'] and description.value==source['description']
        assert title.data_type=='s' and description.data_type=='s'
        assert title.alignment.wrap_text and description.alignment.wrap_text
        assert description.alignment.vertical=='top'
        assert all(sheet.cell(row,c).data_type!='f' for c in range(1,4))
    assert sheet.column_dimensions['C'].width>=80
    results.append({'file':f'{name}.xlsx','records':len(expected),'columns':3,'textExact':True,'datesTyped':True,'wrapped':True})
(root/'independent-reader-result.json').write_text(json.dumps({'reader':'openpyxl '+openpyxl.__version__,'results':results},ensure_ascii=False,indent=2)+'\n')
print('Planos: XLSX único e múltiplo validados; datas, texto exato, Unicode, quebras, fórmulas literais, dimensões e estilos.')
