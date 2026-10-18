"""Validação independente: openpyxl + csv padrão, sem chamar o escritor Kotlin."""
from pathlib import Path
import csv,io,json,sys,math,zipfile
import openpyxl
root=Path(sys.argv[1] if len(sys.argv)>1 else 'app/build/exports-validation')
cases=json.loads((root/'expected.json').read_text());validated=[]
def equal(a,b):
    if a in (None,'') and b in (None,''):return True
    if isinstance(a,(float,int)) and isinstance(b,(float,int)):return math.isclose(a,b,rel_tol=1e-12,abs_tol=1e-12)
    return a==b
for case in cases:
    name=case['name'];expected=case['rows'];width=case['columns'];xlsx=root/(name+'.xlsx')
    with zipfile.ZipFile(xlsx) as z:
        assert z.testzip() is None and 'xl/worksheets/sheet1.xml' in z.namelist()
    book=openpyxl.load_workbook(xlsx,data_only=False);sheet=book.active
    assert sheet.max_row==max(1,len(expected)) and sheet.max_column==width,(name,sheet.max_row,sheet.max_column)
    assert not sheet.protection.sheet
    for r,row in enumerate(expected,1):
        for c,value in enumerate(row,1):
            cell=sheet.cell(r,c);assert equal(cell.value,value),(name,cell.coordinate,cell.value,value)
            if isinstance(value,(int,float)):
                assert cell.data_type=='n'
                if cell.number_format!='General':assert '0.00' in cell.number_format
            assert cell.data_type!='f'
    csv_bytes=(root/(name+'.csv')).read_bytes();assert csv_bytes.startswith(b'\xef\xbb\xbf')
    for suffix,delimiter,rows in [('.csv',';',expected),('.tsv','\t',expected),('-notes.tsv','\t',case['notes']),('-all-notes.tsv','\t',case['allNotes'])]:
        text=(root/(name+suffix)).read_bytes().decode('utf-8-sig')
        actual=list(csv.reader(io.StringIO(text,newline=''),delimiter=delimiter))
        assert len(actual)==len(rows),(name,suffix,len(actual),len(rows))
        for output,source in zip(actual,rows):
            assert len(output)==len(source),(name,suffix,output,source)
            for value,original in zip(output,source):
                if original is None:assert value==''
                elif isinstance(original,(int,float)):assert math.isclose(float(value.replace(',','.')),original,rel_tol=1e-12,abs_tol=1e-12)
                else:assert value==original,(name,suffix,value,original)
    validated.append({'case':name,'rows':len(expected),'columns':width,'xlsx':'valid','csv':'valid','tsv':'valid'})
(root/'independent-reader-result.json').write_text(json.dumps({'reader':'openpyxl '+openpyxl.__version__,'cases':len(validated),'results':validated},ensure_ascii=False,indent=2)+'\n')
print(f'{len(validated)} XLSX reais, {len(validated)} CSV e {len(validated)*3} TSV validados com leitor independente; células, tipos, dimensões, Unicode e espaços conferidos.')
